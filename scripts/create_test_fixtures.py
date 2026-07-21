from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "tests" / "fixtures"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 98, 108)


def set_font(run, size: float, *, bold: bool = False, color: RGBColor | None = None) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def create_docx(path: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_font(header.add_run("RAG INGESTION FIXTURE  |  INTERNAL"), 9, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("Format validation fixture - no sensitive information"), 8.5, color=MUTED)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("知识库接入操作指南"), 23, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    set_font(subtitle.add_run("DOCX 格式端到端验收样本"), 13, color=MUTED)

    for label, value in (
        ("文档类型", "DOCX"),
        ("授权部门", "全公司"),
        ("测试标识", "DOCX-READY-2026"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        set_font(paragraph.add_run(f"{label}："), 11, bold=True)
        set_font(paragraph.add_run(value), 11)

    doc.add_heading("上传前检查", level=1)
    for text in (
        "确认文件不包含密码、密钥、个人身份号码等敏感信息。",
        "选择正确的部门权限，并使用清晰可检索的文档标题。",
        "文件大小不得超过 20 MB。",
    ):
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("处理流程", level=1)
    for text in (
        "管理员上传文件，系统计算 SHA-256 并创建索引任务。",
        "Worker 解析段落、切块并生成本地 BGE 向量。",
        "索引完成后状态变为 ready，员工按部门权限检索。",
    ):
        doc.add_paragraph(text, style="List Number")

    doc.add_heading("验收口径", level=1)
    doc.add_paragraph(
        "测试口令 DOCX-READY-2026 必须能够从解析文本中找到。重复上传应返回冲突；删除文档后，文档记录和向量片段必须同步消失。"
    )
    doc.save(path)


def create_pdf(path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter)
    width, height = letter
    for page, (heading, lines) in enumerate(
        (
            (
                "Procurement Approval Flow",
                [
                    "Validation token: PDF-READY-2026",
                    "Requests up to CNY 5,000 require manager approval.",
                    "Requests above CNY 5,000 also require the finance controller.",
                    "Requests above CNY 20,000 require three comparable quotations.",
                ],
            ),
            (
                "Evidence and Retention",
                [
                    "Attach the invoice, approval record, and supplier quotation.",
                    "Procurement evidence is retained for 36 months.",
                    "The requester must explain any emergency single-source purchase.",
                ],
            ),
        ),
        start=1,
    ):
        pdf.setFillColorRGB(0.12, 0.25, 0.42)
        pdf.setFont("Helvetica-Bold", 20)
        pdf.drawString(72, height - 90, heading)
        pdf.setFillColorRGB(0.35, 0.39, 0.44)
        pdf.setFont("Helvetica", 9)
        pdf.drawRightString(width - 72, height - 45, f"RAG PDF fixture | Page {page} of 2")
        pdf.setFillColorRGB(0, 0, 0)
        pdf.setFont("Helvetica", 11)
        y = height - 145
        for line in lines:
            pdf.drawString(88, y, f"- {line}")
            y -= 30
        pdf.showPage()
    pdf.save()


def create_scan_only_pdf(path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter)
    pdf.setLineWidth(2)
    pdf.rect(90, 160, 430, 470)
    for y in range(580, 220, -45):
        pdf.line(125, y, 485, y)
    pdf.circle(305, 690, 28)
    pdf.save()


def validate(docx_path: Path, pdf_path: Path, scan_path: Path) -> None:
    docx_text = "\n".join(paragraph.text for paragraph in Document(docx_path).paragraphs)
    assert "DOCX-READY-2026" in docx_text
    pdf_pages = PdfReader(str(pdf_path)).pages
    assert len(pdf_pages) == 2
    assert "PDF-READY-2026" in "\n".join(page.extract_text() or "" for page in pdf_pages)
    assert not "".join(page.extract_text() or "" for page in PdfReader(str(scan_path)).pages).strip()


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    docx_path = OUTPUT / "知识库接入操作指南.docx"
    pdf_path = OUTPUT / "procurement-approval-flow.pdf"
    scan_path = OUTPUT / "scan-only-no-text.pdf"
    create_docx(docx_path)
    create_pdf(pdf_path)
    create_scan_only_pdf(scan_path)
    (OUTPUT / "format-validation.md").write_text(
        "# Markdown 格式验收\n\n测试标识：MD-READY-2026。\n\n该文件用于验证 Markdown 上传、索引、去重与删除流程。\n",
        encoding="utf-8",
    )
    validate(docx_path, pdf_path, scan_path)
    print(f"created fixtures in {OUTPUT}")


if __name__ == "__main__":
    main()
